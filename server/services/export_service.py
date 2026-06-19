"""Export translated content to common formats."""

from __future__ import annotations

import io
import json
from dataclasses import dataclass

from server.schemas.export_schema import ExportRequest


@dataclass
class ExportedFile:
    filename: str
    media_type: str
    content: bytes


class ExportService:
    def export(self, req: ExportRequest) -> ExportedFile:
        if req.format == "json":
            return self._json(req)
        if req.format == "srt":
            return self._srt(req)
        if req.format == "docx":
            return self._docx(req)
        return self._txt(req)

    @staticmethod
    def _txt(req: ExportRequest) -> ExportedFile:
        parts = [req.title.strip() or "SmartGen Export", ""]
        if req.source_text:
            parts.extend(["Source:", req.source_text.strip(), ""])
        if req.translated_text:
            parts.extend(["Translation:", req.translated_text.strip(), ""])
        for seg in req.segments:
            parts.append(f"[{seg.start:.2f} - {seg.end:.2f}] {seg.text}")
            if seg.translation:
                parts.append(seg.translation)
            parts.append("")
        return ExportedFile("smartgen-export.txt", "text/plain; charset=utf-8", "\n".join(parts).encode("utf-8"))

    @staticmethod
    def _json(req: ExportRequest) -> ExportedFile:
        content = json.dumps(req.model_dump(), ensure_ascii=False, indent=2).encode("utf-8")
        return ExportedFile("smartgen-export.json", "application/json; charset=utf-8", content)

    @staticmethod
    def _srt(req: ExportRequest) -> ExportedFile:
        if not req.segments:
            raise ValueError("segments are required for srt export")
        blocks = []
        for i, seg in enumerate(req.segments, 1):
            text = seg.translation or seg.text
            blocks.append(f"{i}\n{ExportService._srt_time(seg.start)} --> {ExportService._srt_time(seg.end)}\n{text}\n")
        return ExportedFile("smartgen-export.srt", "application/x-subrip; charset=utf-8", "\n".join(blocks).encode("utf-8"))

    @staticmethod
    def _docx(req: ExportRequest) -> ExportedFile:
        try:
            from docx import Document
        except Exception as exc:
            raise RuntimeError("DOCX export requires optional dependency: pip install python-docx") from exc

        doc = Document()
        doc.add_heading(req.title.strip() or "SmartGen Export", 0)
        if req.source_text:
            doc.add_heading("Source", level=1)
            doc.add_paragraph(req.source_text)
        if req.translated_text:
            doc.add_heading("Translation", level=1)
            doc.add_paragraph(req.translated_text)
        if req.segments:
            doc.add_heading("Segments", level=1)
            for seg in req.segments:
                doc.add_paragraph(f"{seg.start:.2f} - {seg.end:.2f}", style=None)
                doc.add_paragraph(seg.text)
                if seg.translation:
                    doc.add_paragraph(seg.translation)
        buf = io.BytesIO()
        doc.save(buf)
        return ExportedFile(
            "smartgen-export.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            buf.getvalue(),
        )

    @staticmethod
    def _srt_time(seconds: float) -> str:
        ms_total = max(0, int(round(seconds * 1000)))
        hours, rem = divmod(ms_total, 3600_000)
        minutes, rem = divmod(rem, 60_000)
        secs, ms = divmod(rem, 1000)
        return f"{hours:02}:{minutes:02}:{secs:02},{ms:03}"
